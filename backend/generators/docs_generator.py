"""Docs generator - turns a prompt into a real, downloadable .docx file.

Two-step process: (1) ask the LLM to draft structured content as JSON
(title + sections, each with a heading and body paragraphs), (2) render
that JSON into an actual Word document with python-docx. Splitting it this
way means the LLM never has to know anything about docx formatting - it
just writes content, and this module handles turning it into a real file.
"""
from __future__ import annotations

import json
import os
import re
import uuid
from typing import Any, Dict, List

from docx import Document
from docx.shared import Pt


DRAFT_SYSTEM_PROMPT = (
    "You are a document drafting assistant. Given a topic, produce a "
    "well-organized document as JSON with this exact shape:\n"
    '{"title": "...", "sections": [{"heading": "...", "paragraphs": ["...", "..."]}]}\n'
    "Return ONLY the JSON, no other text, no markdown code fences."
)


def _extract_json(text: str) -> Dict[str, Any]:
    """LLMs sometimes wrap JSON in prose or code fences despite instructions
    not to - strip that defensively rather than failing the whole request."""
    text = text.strip()
    fence_match = re.search(r"```(?:json)?\s*(\{.*\})\s*```", text, re.DOTALL)
    if fence_match:
        text = fence_match.group(1)
    else:
        brace_match = re.search(r"\{.*\}", text, re.DOTALL)
        if brace_match:
            text = brace_match.group(0)
    return json.loads(text)


def draft_document_content(llm, topic: str) -> Dict[str, Any]:
    """Ask the LLM to draft structured content for the topic. Raises
    ValueError with the raw response included if the model didn't return
    parseable JSON, rather than silently producing an empty/broken doc."""
    response = llm.chat(
        [{"role": "user", "content": f"Write a document about: {topic}"}],
        system_prompt=DRAFT_SYSTEM_PROMPT,
        temperature=0.5,
    )
    try:
        return _extract_json(response.text)
    except (json.JSONDecodeError, AttributeError) as exc:
        raise ValueError(
            f"Model didn't return valid JSON for document content: {exc}. "
            f"Raw response started with: {response.text[:200]!r}"
        ) from exc


def build_docx(content: Dict[str, Any], output_dir: str) -> str:
    """Render drafted content into a real .docx file. Returns the file path."""
    doc = Document()

    title = content.get("title", "Untitled Document")
    doc.add_heading(title, level=0)

    for section in content.get("sections", []):
        heading = section.get("heading", "")
        if heading:
            doc.add_heading(heading, level=1)
        for para in section.get("paragraphs", []):
            p = doc.add_paragraph(para)
            p.style.font.size = Pt(11)

    os.makedirs(output_dir, exist_ok=True)
    filename = f"doc_{uuid.uuid4().hex[:10]}.docx"
    path = os.path.join(output_dir, filename)
    doc.save(path)
    return path


def generate_docx(llm, topic: str, output_dir: str) -> str:
    """End-to-end: draft content with the LLM, render it to a real .docx
    file, return the file path."""
    content = draft_document_content(llm, topic)
    return build_docx(content, output_dir)
