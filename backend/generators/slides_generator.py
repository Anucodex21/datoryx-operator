"""Slides generator - turns a prompt into a real, downloadable .pptx file.

Same two-step pattern as docs_generator.py: LLM drafts structured slide
content as JSON, this module renders it into an actual PowerPoint file.
"""
from __future__ import annotations

import json
import os
import re
import uuid
from typing import Any, Dict

from pptx import Presentation
from pptx.util import Inches, Pt


DRAFT_SYSTEM_PROMPT = (
    "You are a presentation drafting assistant. Given a topic, produce slide "
    "content as JSON with this exact shape:\n"
    '{"title": "...", "slides": [{"heading": "...", "bullets": ["...", "..."]}]}\n'
    "Aim for 5-8 slides. Return ONLY the JSON, no other text, no markdown code fences."
)


def _extract_json(text: str) -> Dict[str, Any]:
    text = text.strip()
    fence_match = re.search(r"```(?:json)?\s*(\{.*\})\s*```", text, re.DOTALL)
    if fence_match:
        text = fence_match.group(1)
    else:
        brace_match = re.search(r"\{.*\}", text, re.DOTALL)
        if brace_match:
            text = brace_match.group(0)
    return json.loads(text)


def draft_slide_content(llm, topic: str) -> Dict[str, Any]:
    response = llm.chat(
        [{"role": "user", "content": f"Create a slide deck about: {topic}"}],
        system_prompt=DRAFT_SYSTEM_PROMPT,
        temperature=0.5,
    )
    try:
        return _extract_json(response.text)
    except (json.JSONDecodeError, AttributeError) as exc:
        raise ValueError(
            f"Model didn't return valid JSON for slide content: {exc}. "
            f"Raw response started with: {response.text[:200]!r}"
        ) from exc


def build_pptx(content: Dict[str, Any], output_dir: str) -> str:
    """Render drafted content into a real .pptx file. Returns the file path."""
    prs = Presentation()

    title_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(title_layout)
    slide.shapes.title.text = content.get("title", "Untitled Deck")

    bullet_layout = prs.slide_layouts[1]
    for slide_data in content.get("slides", []):
        slide = prs.slides.add_slide(bullet_layout)
        slide.shapes.title.text = slide_data.get("heading", "")
        body = slide.placeholders[1].text_frame
        bullets = slide_data.get("bullets", [])
        if bullets:
            body.text = bullets[0]
            for bullet in bullets[1:]:
                p = body.add_paragraph()
                p.text = bullet

    os.makedirs(output_dir, exist_ok=True)
    filename = f"slides_{uuid.uuid4().hex[:10]}.pptx"
    path = os.path.join(output_dir, filename)
    prs.save(path)
    return path


def generate_pptx(llm, topic: str, output_dir: str) -> str:
    content = draft_slide_content(llm, topic)
    return build_pptx(content, output_dir)
